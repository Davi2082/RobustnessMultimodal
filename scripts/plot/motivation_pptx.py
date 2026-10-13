"""Editable PowerPoint version of the Introduction's motivation figure.

Geometry is imported from scripts/plot/plot_motivation.py, so both versions
always match. Every element is a native, named shape (Selection Pane); the
decision boundary is one freeform curve (Edit Points to reshape it).
To use it in the paper: File > Export > PDF, then crop.

Usage (needs python-pptx, numpy, matplotlib):
    python3 scripts/plot/motivation_pptx.py  ->  figures/motivation/motivation.pptx
"""

import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

sys.path.insert(0, os.path.dirname(__file__))
import plot_motivation as G  # noqa: E402

def hexc(c): return c.lstrip("#").upper()
TEXT_C, IMAGE_C, JOINT_C, FLIP_C, INK = map(hexc, (G.TEXT_C, G.IMAGE_C, G.JOINT_C, G.FAKE_C, G.INK))
WHITE, LABEL_GREY = "FFFFFF", "666666"
ADV_C = hexc(G.ADV_C)
FONT = "Times New Roman"

S = 0.31                           # inches per plot unit
LEFT, TOP, BOTTOM = 0.45, 0.32, 0.65
GAP_X = 0.35                      # space between the two panels
OFFSET = [0.0]                    # x offset of the panel being drawn
SLIDE_W, SLIDE_H = 2 * (LEFT + G.XMAX * S) + GAP_X, TOP + G.YMAX * S + BOTTOM


def X(x): return OFFSET[0] + LEFT + x * S
def Y(y): return TOP + (G.YMAX - y) * S


def alpha(shp, opacity):
    clr = shp._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    clr.append(clr.makeelement(qn("a:alpha"), {"val": str(int(opacity * 100000))}))


def poly(slide, name, pts, closed, fill=None, fill_opacity=1.0, line=None, line_w=1.0, dash=False):
    pts = [(X(x), Y(y)) for x, y in pts]
    fb = slide.shapes.build_freeform(Inches(pts[0][0]), Inches(pts[0][1]), scale=1.0)
    fb.add_line_segments([(Inches(x), Inches(y)) for x, y in pts[1:]], close=closed)
    s = fb.convert_to_shape()
    s.name = name
    s.shadow.inherit = False
    if fill:
        s.fill.solid(); s.fill.fore_color.rgb = RGBColor.from_string(fill)
        if fill_opacity < 1:
            alpha(s, fill_opacity)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = RGBColor.from_string(line); s.line.width = Pt(line_w)
        if dash:
            s.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    else:
        s.line.fill.background()
    return s


def segment(slide, name, a, b, color, width=1.0, dash=False, arrow=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(X(a[0])), Inches(Y(a[1])),
                                   Inches(X(b[0])), Inches(Y(b[1])))
    c.name = name
    c.line.color.rgb = RGBColor.from_string(color)
    c.line.width = Pt(width)
    if dash:
        c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if arrow:
        ln = c.line._get_or_add_ln()
        ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return c


def dot(slide, name, pt, color, size=0.09, shape=MSO_SHAPE.OVAL):
    s = slide.shapes.add_shape(shape, Inches(X(pt[0]) - size / 2), Inches(Y(pt[1]) - size / 2),
                               Inches(size), Inches(size))
    s.name = name
    s.shadow.inherit = False
    s.fill.solid(); s.fill.fore_color.rgb = RGBColor.from_string(color)
    s.line.fill.background()


def text(slide, name, t, x, y, w, h, color, size=11, italic=False, bold=False,
         align=PP_ALIGN.CENTER, rotation=0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.name = name
    tb.rotation = rotation
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = False
    para = tf.paragraphs[0]
    para.alignment = align
    r = para.add_run()
    r.text = t
    r.font.name, r.font.size, r.font.italic, r.font.bold = FONT, Pt(size), italic, bold
    r.font.color.rgb = RGBColor.from_string(color)


def build(slide, panel):
    curve = G.boundary()
    tag = "a" if panel == "unimodal" else "b"
    poly(slide, f"{tag}: " + "real region", curve + [(G.XMAX, 0), (G.XMAX, G.YMAX), (0, G.YMAX)], True,
         fill=FLIP_C, fill_opacity=0.10)
    poly(slide, f"{tag}: " + "decision boundary", curve, False, line=INK, line_w=1.5)
    segment(slide, f"{tag}: " + "x axis", (0, 0), (G.XMAX, 0), INK)
    segment(slide, f"{tag}: " + "y axis", (0, 0), (0, G.YMAX), INK)
    text(slide, f"{tag}: " + "label real", "real", X(G.XMAX) - 0.8, Y(G.YMAX) + 0.04, 0.75, 0.2, FLIP_C,
         align=PP_ALIGN.RIGHT)
    text(slide, f"{tag}: " + "label fake", "fake", X(0.15), Y(0.15) - 0.2, 0.7, 0.2, LABEL_GREY, align=PP_ALIGN.LEFT)
    text(slide, f"{tag}: " + "x-axis title", "image", X(0), Y(0) + 0.1, G.XMAX * S, 0.22, INK)
    text(slide, f"{tag}: " + "y-axis title", "text", X(0) - 0.4, Y(G.YMAX / 2) - 0.11, 0.4, 0.22, INK, rotation=270)

    for side, pts in G.BACKGROUND.items():
        for j, q in enumerate(pts):
            shp = MSO_SHAPE.ISOSCELES_TRIANGLE if side == "fake" else MSO_SHAPE.OVAL
            sz = 0.1
            m = slide.shapes.add_shape(shp, Inches(X(q[0]) - sz / 2), Inches(Y(q[1]) - sz / 2), Inches(sz), Inches(sz))
            m.name = f"{tag} unattacked {side} {j + 1}"
            m.shadow.inherit = False
            m.fill.solid(); m.fill.fore_color.rgb = RGBColor.from_string("9A9A9A"); m.line.fill.background()

    for name, p in G.SAMPLES.items():
        ends = G.ATTACKS[name]
        robust = name in G.ROBUST
        if panel == "unimodal":
            segment(slide, f"{tag} {name}: image attack", p, ends["image"], IMAGE_C, 1.25, arrow=True)
            segment(slide, f"{tag} {name}: text attack", p, ends["text"], TEXT_C, 1.25, arrow=True)
            for m in G.FLIPS[name]:
                dot(slide, f"{tag} {name}: adversarial ({m}), classified real", ends[m], ADV_C, 0.11,
                    MSO_SHAPE.ISOSCELES_TRIANGLE)
        elif robust:
            c = ends["union"]
            for k, (u, v) in enumerate([(p, ends["image"]), (ends["image"], c), (c, ends["text"]), (ends["text"], p)]):
                segment(slide, f"{tag} {name}: box side {k + 1}", u, v, "999999", 0.5, dash=True)
            segment(slide, f"{tag} {name}: union", p, c, JOINT_C, 1.0, dash=True, arrow=True)
            segment(slide, f"{tag} {name}: joint", p, ends["joint"], JOINT_C, 1.5, arrow=True)
            dot(slide, f"{tag} {name}: adversarial (joint), classified real", ends["joint"], ADV_C, 0.11,
                MSO_SHAPE.ISOSCELES_TRIANGLE)
        dot(slide, f"{tag} {name}: sample", p, INK, 0.1, MSO_SHAPE.ISOSCELES_TRIANGLE)
        low = p[1] < 1.0
        text(slide, f"{tag} {name}: label", name, X(p[0]) - (0.3 if low else 0.24), Y(p[1]) + (-0.1 if low else 0.04), 0.16, 0.2,
             INK, size=12, italic=True, align=PP_ALIGN.RIGHT)
    title = "(a) Unimodal attacks" if panel == "unimodal" else "(b) Union and joint attacks"
    text(slide, f"{tag}: title", title, X(0), 0.04, G.XMAX * S, 0.24, INK, size=12)


def legend(slide):
    """One row: attack types, then sample types."""
    y = Y(0) + 0.48
    items = [("text", TEXT_C, "line"), ("image", IMAGE_C, "line"), ("union", JOINT_C, "dash"),
             ("joint", JOINT_C, "line"), ("real", "9A9A9A", MSO_SHAPE.OVAL),
             ("fake", INK, MSO_SHAPE.ISOSCELES_TRIANGLE), ("adversarial fake", ADV_C, MSO_SHAPE.ISOSCELES_TRIANGLE)]
    widths = [0.28 + 0.06 + 0.075 * len(label) + 0.18 for label, _, _ in items]
    x = (SLIDE_W - sum(widths)) / 2
    for (label, color, kind), w in zip(items, widths):
        if kind in ("line", "dash"):
            c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x + 0.28), Inches(y))
            c.name = f"legend: {label} line"
            c.line.color.rgb = RGBColor.from_string(color); c.line.width = Pt(1.25)
            if kind == "dash":
                c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        else:
            m = slide.shapes.add_shape(kind, Inches(x + 0.09), Inches(y - 0.05), Inches(0.1), Inches(0.1))
            m.name = f"legend: {label} marker"
            m.shadow.inherit = False
            m.fill.solid(); m.fill.fore_color.rgb = RGBColor.from_string(color); m.line.fill.background()
        text(slide, f"legend: {label}", label, x + 0.34, y - 0.1, 0.075 * len(label) + 0.1, 0.2, INK,
             size=9, align=PP_ALIGN.LEFT)
        x += w


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    for k, panel in enumerate(("unimodal", "joint")):
        OFFSET[0] = k * (LEFT + G.XMAX * S + GAP_X)
        build(slide, panel)
    OFFSET[0] = 0.0
    legend(slide)
    out = "figures/motivation/motivation.pptx"
    os.makedirs(os.path.dirname(out), exist_ok=True)
    prs.save(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    main()
